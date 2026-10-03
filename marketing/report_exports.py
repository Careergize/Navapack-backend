"""Download renderers for the existing five-section CRM report."""
import math
import re
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from django.http import FileResponse
from rest_framework import serializers
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from .reports import build_report


class ISODateField(serializers.DateField):
    def to_internal_value(self, value):
        if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
            self.fail('invalid', format='YYYY-MM-DD')
        return super().to_internal_value(value)


class ReportExportParameters(serializers.Serializer):
    from_date = ISODateField(input_formats=['%Y-%m-%d'])
    to_date = ISODateField(input_formats=['%Y-%m-%d'])
    salesperson = serializers.IntegerField(required=False, min_value=1)
    limit = serializers.IntegerField(required=False, default=10, min_value=1, max_value=100)

    def validate(self, data):
        if data['from_date'] > data['to_date']:
            raise serializers.ValidationError({'error': 'From date cannot be greater than To date.'})
        return data


# Columns follow the uploaded report's five sections. Types drive both renderers.
SECTIONS = [
    ('top_opportunities', 'Top Opportunities', 'TOP ACTIVE SALES OPPORTUNITIES', [
        ('customer', 'Customer', 'text', 14), ('salesperson', 'Salesperson', 'text', 12),
        ('product', 'Product', 'text', 16), ('potential_value_ugx', 'Potential Value (UGX)', 'money', 11),
        ('sales_stage', 'Current Sales Stage', 'text', 14), ('next_action', 'Next Action', 'text', 21),
        ('next_followup_date', 'Follow-up Date', 'date', 12)]),
    ('market_intelligence', 'Market Intelligence', 'COMPETITOR / RAW MATERIAL / PRICE INTELLIGENCE', [
        ('date', 'Date', 'date', 10), ('salesperson', 'Salesperson', 'text', 12),
        ('customer', 'Customer', 'text', 13), ('area', 'Area / Location', 'text', 13),
        ('intelligence', 'Market / Competitor Information', 'text', 29),
        ('recommended_action', 'Recommended Action', 'text', 23)]),
    ('management_delays', 'Management Delays', 'MANAGEMENT / PRODUCTION / SAMPLE DELAYS TO RESOLVE', [
        ('date', 'Date', 'date', 10), ('customer', 'Customer / Company', 'text', 13),
        ('salesperson', 'Salesperson', 'text', 12), ('support_required', 'Support / Issue Required', 'text', 23),
        ('responsible', 'Responsible Person / Dept.', 'text', 15),
        ('required_by_date', 'Required By / Deadline', 'date', 12), ('status', 'Status', 'text', 15)]),
    ('orders_lost', 'Orders Lost', 'ORDERS LOST & ROOT CAUSE ANALYSIS', [
        ('customer', 'Customer / Company', 'text', 13), ('salesperson', 'Salesperson', 'text', 12),
        ('product', 'Product / Service', 'text', 16), ('potential_value_ugx', 'Potential Value (UGX)', 'money', 11),
        ('reason_lost', 'Reason Lost', 'text', 17), ('competitor', 'Competitor / Supplier', 'text', 13),
        ('corrective_action', 'Corrective Action', 'text', 18)]),
]


def report_tables(report):
    performance = report['sales_performance']
    people = performance['salespersons']
    columns = [('metric', 'KPI / Metric', 'text', 30)] + [
        (f'person_{i}', name, 'number', 14) for i, name in enumerate(people)] + [
        ('total_team', 'TOTAL TEAM', 'number', 16)]
    rows = [{'metric': item['metric'], 'total_team': item['total_team'],
             '_currency': item['is_currency'], **{
                 f'person_{i}': item['per_salesperson'][name] for i, name in enumerate(people)}}
            for item in performance['rows']]
    return [('Sales Performance', 'SALES PERFORMANCE', columns, rows)] + [
        (sheet, title, columns, report[key]) for key, sheet, title, columns in SECTIONS]


def clean_text(value):
    # Excel XML cannot store most ASCII control characters.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))


def excel_report(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, title, columns, rows in report_tables(report):
        sheet = workbook.create_sheet(name)
        count = len(columns)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
        sheet.cell(1, 1, 'NAVAPACK CRM REPORT').font = Font(size=20, color='243447', bold=True)
        sheet.cell(1, 1).fill = PatternFill('solid', fgColor='F8CBAD')
        sheet.row_dimensions[1].height = 32
        sheet.cell(2, 1, 'From Date:')
        sheet.cell(2, 2, report['start_date']).number_format = 'dd mmm yyyy'
        sheet.cell(3, 1, 'To Date:')
        sheet.cell(3, 2, report['end_date']).number_format = 'dd mmm yyyy'
        sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=count)
        sheet.cell(5, 1, title).font = Font(bold=True, size=12, color='243447')
        sheet.row_dimensions[5].height = 25
        sheet.append([column[1] for column in columns])  # row 6
        for cell in sheet[6]:
            cell.data_type = 's'
            cell.font = Font(bold=True, color='243447')
            cell.fill = PatternFill('solid', fgColor='9BC2E6' if name == 'Sales Performance' else '00B0F0')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[6].height = 36
        for row_number, row in enumerate(rows, 7):
            for col_number, (key, label, kind, weight) in enumerate(columns, 1):
                value = row.get(key)
                cell = sheet.cell(row_number, col_number)
                if isinstance(value, str):
                    cell.value = clean_text(value)
                    cell.data_type = 's'  # customer text beginning '=' must never become a formula
                else:
                    cell.value = value
                cell.font = Font(size=11, bold=(key == 'total_team'))
                cell.alignment = Alignment(wrap_text=kind not in ('money', 'number'), vertical='top',
                    horizontal='right' if kind in ('money', 'number') else 'left')
                cell.border = Border(bottom=Side(style='hair', color='DCE4EC'))
                if row_number % 2:
                    cell.fill = PatternFill('solid', fgColor='F2F6FA')
                if kind == 'date':
                    cell.number_format = 'dd mmm yyyy'
                elif kind == 'money' or (kind == 'number' and row.get('_currency')):
                    cell.number_format = '"UGX "#,##0.00;[Red]("UGX "#,##0.00)'
                elif kind == 'number':
                    cell.number_format = '#,##0'
        if not rows:
            sheet.cell(7, 1, 'No records in the selected period.')
        for index, (_, label, kind, _) in enumerate(columns, 1):
            texts = [label]
            for r in range(7, sheet.max_row + 1):
                cell = sheet.cell(r, index)
                value = cell.value
                if isinstance(value, date):
                    text = value.strftime('%d %b %Y')
                elif isinstance(value, (int, float)):
                    text = f'UGX {value:,.2f}' if 'UGX' in cell.number_format else f'{value:,}'
                else:
                    text = str(value or '')
                texts.append(text)
            currency_column = any('UGX' in sheet.cell(r, index).number_format
                                  for r in range(7, sheet.max_row + 1))
            minimum = 24 if currency_column else (16 if kind == 'date' else 18)
            width = max(minimum, min(46, max(map(len, texts), default=14) + 4))
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[6].height = max(36, max(math.ceil(len(column[1]) /
            (sheet.column_dimensions[get_column_letter(i)].width - 2)) * 15 + 8
            for i, column in enumerate(columns, 1)))
        # Approximate wrapped line heights, allowing long source notes without hidden text.
        for row_number in range(7, sheet.max_row + 1):
            lines = max(sum(max(1, math.ceil(len(line) / max(1,
                sheet.column_dimensions[get_column_letter(i)].width - 4)))
                for line in str(sheet.cell(row_number, i).value or '').split('\n'))
                for i in range(1, count + 1))
            sheet.row_dimensions[row_number].height = min(409, max(30, lines * 18 + 8))
        sheet.freeze_panes = 'B7' if name == 'Sales Performance' else 'A7'
        sheet.auto_filter.ref = f'A6:{get_column_letter(count)}{max(6, sheet.max_row)}'
        sheet.print_title_rows = '1:6'
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_view.showGridLines = False
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def pdf_report(report):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    # ReportLab ships these fonts, so deployment does not rely on Windows font files.
    font_dir = Path(reportlab.__file__).parent / 'fonts'
    if 'NavapackVera' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('NavapackVera', str(font_dir / 'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('NavapackVeraBold', str(font_dir / 'VeraBd.ttf')))
    normal = ParagraphStyle('Cell', fontName='NavapackVera', fontSize=8, leading=11,
                            textColor=colors.HexColor('#243447'), splitLongWords=True)
    header = ParagraphStyle('Header', parent=normal, fontName='NavapackVeraBold')
    right = ParagraphStyle('Number', parent=normal, alignment=TA_RIGHT)
    heading = ParagraphStyle('Heading', parent=header, fontSize=14, leading=19, spaceAfter=12)
    brand = ParagraphStyle('Brand', parent=heading, fontSize=24, leading=28, spaceAfter=6)
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=30, rightMargin=30,
        topMargin=32, bottomMargin=36, title='NAVAPACK CRM REPORT', author='NAVAPACK')
    period = f'{report["start_date"]:%d %B %Y} - {report["end_date"]:%d %B %Y}'
    story = []

    def paragraph(value, style=normal):
        return Paragraph(escape(clean_text(value)).replace('\n', '<br/>'), style)

    def add_table(columns, rows):
        values = [[paragraph(column[1], header) for column in columns]]
        for row in rows:
            cells = []
            for key, _, kind, _ in columns:
                value = row.get(key)
                if value is None or value == '':
                    value = '-'
                elif isinstance(value, date):
                    value = value.strftime('%d %b %Y')
                elif kind == 'money' or (kind == 'number' and row.get('_currency')):
                    value = f'UGX {value:,.2f}'
                elif kind == 'number':
                    value = f'{value:,.0f}'
                cells.append(paragraph(value, right if kind in ('number', 'money') else normal))
            values.append(cells)
        if not rows:
            story.append(paragraph('No records in the selected period.'))
            return
        weights = sum(column[3] for column in columns)
        table = LongTable(values, colWidths=[document.width * column[3] / weights for column in columns],
                          repeatRows=1, splitByRow=1, splitInRow=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#9BC2E6')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F6FA')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ('LINEBELOW', (0, 0), (-1, 0), 0.7, colors.HexColor('#749ABC')),
            ('LINEBELOW', (0, 1), (-1, -1), 0.3, colors.HexColor('#DCE4EC')),
        ]))
        story.append(table)

    for index, (_, title, columns, rows) in enumerate(report_tables(report)):
        if index:
            story.append(PageBreak())
        story.extend([paragraph('NAVAPACK', brand), paragraph('CRM REPORT', heading),
                      paragraph(f'Report Period: {period}'), Spacer(1, 16), paragraph(title, heading)])
        if index == 0:
            # Continue the matrix in readable groups rather than shrinking unlimited rep columns.
            reps = columns[1:-1]
            groups = [reps[i:i + 5] for i in range(0, len(reps), 5)] or [[]]
            for group_number, group in enumerate(groups):
                if group_number:
                    story.extend([Spacer(1, 16), paragraph('Sales performance continued', heading)])
                add_table([columns[0], *group, columns[-1]], rows)
        else:
            add_table(columns, rows)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('NavapackVera', 8)
        canvas.setFillColor(colors.HexColor('#62788A'))
        canvas.drawString(30, 19, f'NAVAPACK CRM REPORT | {period}')
        canvas.drawRightString(landscape(A4)[0] - 30, 19, f'Page {doc.page}')
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output


class ReportExportAPIView(APIView):
    permission_classes = [IsAuthenticated]
    export_format = None

    def get(self, request):
        params = ReportExportParameters(data=request.query_params)
        if not params.is_valid():
            errors = params.errors
            key = next(iter(errors))
            message = str(errors[key][0])
            if key != 'error':
                message = f'{key}: {message}'
            return Response({'error': message}, status=400)
        data = params.validated_data
        report = build_report(data['from_date'], data['to_date'],
            salesperson_id=data.get('salesperson'), limit=data['limit'],
            touched_only=True, period_only=True)
        if self.export_format == 'xlsx':
            output = excel_report(report)
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:
            output = pdf_report(report)
            content_type = 'application/pdf'
        filename = f'Navapack_Report_{data["from_date"]}_to_{data["to_date"]}.{self.export_format}'
        response = FileResponse(output, as_attachment=True, filename=filename, content_type=content_type)
        response['Cache-Control'] = 'private, no-store'
        return response


class ExcelReportExportAPIView(ReportExportAPIView):
    export_format = 'xlsx'


class PDFReportExportAPIView(ReportExportAPIView):
    export_format = 'pdf'
