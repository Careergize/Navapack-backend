from datetime import date
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from .models import CustomerPipeline, DailyActivity, Salesperson
from .reports import build_report
from .report_exports import excel_report, pdf_report


class ReportExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username='report-user')
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.person = Salesperson.objects.create(name='Alex', user=self.user)
        self.other = Salesperson.objects.create(name='Other')
        self.start, self.end = date(2026, 9, 1), date(2026, 9, 30)
        self.params = {'from_date': '2026-09-01', 'to_date': '2026-09-30'}
        self.deal('Inside customer', date(2026, 9, 15), 150000, next_followup_date=date(2026, 9, 20))
        self.deal('Outside customer', date(2026, 8, 15), 900000, next_followup_date=date(2026, 8, 20))
        self.deal('Lost inside', date(2026, 9, 21), 42000, sales_stage='Order Lost', reason_lost='Price')
        self.deal('Lost outside', date(2026, 8, 21), 99999, sales_stage='Lost')
        self.deal('Won inside', date(2026, 9, 22), 1000, sales_stage='Won', actual_order_value_ugx=25000)
        for day, company in [(self.start, 'Boundary start'), (self.end, 'Boundary end'),
                              (date(2026, 8, 31), 'Excluded before'), (date(2026, 10, 1), 'Excluded after')]:
            DailyActivity.objects.create(date=day, salesperson=self.person, customer_company=company,
                market_competitor_intelligence='Competitor price', management_support_needed='Sample needed',
                responsible_person_dept='Production', required_by_date=day,
                cash_collected_ugx=500, quotation_submitted_value_ugx=1000)

    def deal(self, company, day, value, **extra):
        obj = CustomerPipeline.objects.create(prospect_id=company, salesperson=self.person,
            customer_company=company, date_added=day, last_contact_date=day,
            estimated_value_ugx=value, product_service='Packaging', next_action='Call customer', **extra)
        # auto_now updates on save; set explicitly to emulate stored period records.
        CustomerPipeline.objects.filter(pk=obj.pk).update(stage_last_updated=day)
        return obj

    def download(self, kind, params=None):
        response = self.client.get(reverse('reports-export-' + kind), params or self.params)
        self.assertEqual(response.status_code, 200)
        content = b''.join(response.streaming_content)
        response.close()
        return response, content

    def test_dates_required_valid_and_ordered_in_both_formats(self):
        cases = [{}, {'from_date': '2026-09-01'}, {'to_date': '2026-09-30'},
            {'from_date': '', 'to_date': '2026-09-30'},
            {'from_date': '2026-02-30', 'to_date': '2026-09-30'},
            {'from_date': '2026-9-01', 'to_date': '2026-09-30'},
            {'from_date': '2026-09-01T00:00:00', 'to_date': '2026-09-30'},
            {'from_date': '2026-10-01', 'to_date': '2026-09-30'}]
        for kind in ('excel', 'pdf'):
            for params in cases:
                with self.subTest(kind=kind, params=params):
                    response = self.client.get(reverse('reports-export-' + kind), params)
                    self.assertEqual(response.status_code, 400)
                    self.assertIn('error', response.data)
            self.assertEqual(response.data['error'], 'From date cannot be greater than To date.')

    def test_authentication_required(self):
        self.client.force_authenticate(None)
        for kind in ('excel', 'pdf'):
            self.assertIn(self.client.get(reverse('reports-export-' + kind), self.params).status_code, (401, 403))

    def test_excel_real_file_sections_types_and_range(self):
        response, content = self.download('excel')
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('Navapack_Report_2026-09-01_to_2026-09-30.xlsx', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, ['Sales Performance', 'Top Opportunities',
            'Market Intelligence', 'Management Delays', 'Orders Lost'])
        self.assertEqual(workbook['Top Opportunities']['A7'].value, 'Inside customer')
        self.assertEqual(workbook['Top Opportunities'].max_row, 7)
        self.assertEqual(workbook['Top Opportunities']['D7'].value, 150000)
        self.assertIn('UGX', workbook['Top Opportunities']['D7'].number_format)
        self.assertEqual(workbook['Top Opportunities']['G7'].value.date(), date(2026, 9, 20))
        self.assertEqual(workbook['Market Intelligence'].max_row, 8)
        self.assertEqual(workbook['Management Delays'].max_row, 9)
        self.assertEqual(workbook['Orders Lost']['A7'].value, 'Lost inside')
        for sheet in workbook:
            self.assertTrue(sheet.freeze_panes)
            self.assertTrue(sheet.auto_filter.ref)
            self.assertEqual(sheet['B2'].value.date(), self.start)
        metrics = {row[0]: row[-1] for row in workbook['Sales Performance'].iter_rows(min_row=7, values_only=True)}
        self.assertEqual(metrics['Physical Visits'], 2)
        self.assertEqual(metrics['Quotation Value (UGX)'], 2000)
        self.assertEqual(metrics['Cash Collections (UGX)'], 1000)
        self.assertEqual(metrics['Order Value Won (UGX)'], 25000)
        self.assertEqual(metrics['Active Overdue Follow-ups'], 1)

    def test_pdf_real_file_and_all_sections(self):
        response, content = self.download('pdf')
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(content.startswith(b'%PDF-'))
        reader = PdfReader(BytesIO(content))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        for expected in ('NAVAPACK', 'CRM REPORT', '01 September 2026 - 30 September 2026',
            'SALES PERFORMANCE', 'TOP ACTIVE SALES OPPORTUNITIES',
            'COMPETITOR / RAW MATERIAL / PRICE INTELLIGENCE',
            'MANAGEMENT / PRODUCTION / SAMPLE DELAYS TO RESOLVE', 'ORDERS LOST & ROOT CAUSE ANALYSIS',
            'Inside customer', 'Lost inside', 'Boundary start', 'Boundary end'):
            self.assertIn(expected, text)
        for excluded in ('Outside customer', 'Lost outside', 'Excluded before', 'Excluded after'):
            self.assertNotIn(excluded, text)
        self.assertEqual(len(reader.pages), 5)

    def test_empty_range_and_same_day(self):
        for kind in ('excel', 'pdf'):
            self.download(kind, {'from_date': '2025-01-01', 'to_date': '2025-01-01'})

    def test_salesperson_filter_and_invalid_optional_parameters(self):
        _, content = self.download('excel', dict(self.params, salesperson=self.other.pk))
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook['Sales Performance']['B6'].value, 'Other')
        self.assertEqual(workbook['Top Opportunities']['A7'].value, 'No records in the selected period.')
        for extra in ({'salesperson': 'abc'}, {'salesperson': '-1'}, {'limit': 'abc'}, {'limit': '0'}, {'limit': '101'}):
            self.assertEqual(self.client.get(reverse('reports-export-excel'), dict(self.params, **extra)).status_code, 400)

    def test_exports_reuse_existing_report_builders(self):
        with patch('marketing.report_exports.build_report', wraps=build_report) as builder:
            self.download('excel')
            builder.assert_called_once_with(self.start, self.end, salesperson_id=None, limit=10,
                touched_only=True, period_only=True)

    def test_individual_selection_aliases_filter_every_section_in_both_formats(self):
        for field, value in (('salesperson', self.other.pk),
                             ('salesperson_id', self.other.pk)):
            for kind in ('excel', 'pdf'):
                with self.subTest(field=field, kind=kind):
                    _, content = self.download(kind, dict(self.params, **{field: value}))
                    if kind == 'excel':
                        workbook = load_workbook(BytesIO(content))
                        text = str([list(sheet.values) for sheet in workbook])
                    else:
                        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
                    self.assertIn('Other', text)
                    for excluded in ('Alex', 'Inside customer', 'Lost inside', 'Boundary start'):
                        self.assertNotIn(excluded, text)

    def test_account_user_id_resolves_to_salesperson(self):
        # Account IDs and salesperson IDs are separate namespaces.
        account = get_user_model().objects.create_user(username='selected-account')
        Salesperson.objects.filter(pk=self.person.pk).update(user=account)
        _, content = self.download('excel', dict(self.params, user_id=account.pk))
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook['Sales Performance']['B6'].value, 'Alex')
        self.assertEqual(workbook['Sales Performance']['C6'].value, 'TOTAL TEAM')
        self.assertEqual(workbook['Top Opportunities']['A7'].value, 'Inside customer')

    def test_unknown_and_conflicting_selections_are_rejected(self):
        for extra in ({'salesperson': 999999}, {'salesperson_id': 999999},
                      {'user_id': 999999}, {'user_id': 'abc'},
                      {'salesperson': self.person.pk, 'salesperson_id': self.other.pk}):
            for kind in ('excel', 'pdf'):
                response = self.client.get(reverse('reports-export-' + kind), dict(self.params, **extra))
                self.assertEqual(response.status_code, 400)
                self.assertIn('error', response.data)

    def test_existing_json_defaults_still_include_pipeline_snapshot(self):
        response = self.client.get(reverse('reports'), {'start_date': self.start.isoformat(), 'end_date': self.end.isoformat()})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data['top_opportunities']), 2)
        overdue = next(row for row in response.data['sales_performance']['rows'] if row['key'] == 'overdue')
        self.assertEqual(overdue['total_team'], 2)

    def test_user_text_stays_literal_in_excel_and_pdf(self):
        CustomerPipeline.objects.filter(customer_company='Inside customer').update(
            customer_company='=HYPERLINK("https://example.test")', next_action='<b>Call & confirm</b>\x00')
        _, content = self.download('excel')
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook['Top Opportunities']['A7'].data_type, 's')
        self.assertNotIn('\x00', workbook['Top Opportunities']['F7'].value)
        _, content = self.download('pdf')
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
        self.assertIn('<b>Call & confirm</b>', text)

    def test_many_reps_and_long_notes_paginate_without_truncating(self):
        for i in range(12):
            Salesperson.objects.create(name=f'Rep {i}')
        long_note = 'Long note with business details. ' * 200
        DailyActivity.objects.filter(customer_company='Boundary start').update(market_competitor_intelligence=long_note)
        _, content = self.download('pdf')
        reader = PdfReader(BytesIO(content))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('Rep 11', text)
        self.assertGreater(len(reader.pages), 5)
        # Page headers can interrupt a sentence split across pages. Count each source
        # word separately to verify content is retained rather than flattened adjacency.
        for word in ('Long', 'note', 'business', 'details.'):
            self.assertEqual(text.split().count(word), 200)

    @patch('marketing.reports.timezone.localdate', return_value=date(2026, 10, 18))
    def test_management_delays_include_overdue_pipeline_followups(self, mocked_today):
        self.deal('Completed followup', self.start, 10,
                  next_followup_date=date(2026, 9, 10), followup_status='COMPLETED')
        self.deal('Closed followup', self.start, 10,
                  next_followup_date=date(2026, 9, 10), sales_stage='Order Won')
        other = self.deal('Other rep overdue', self.start, 10,
                          next_followup_date=date(2026, 9, 11))
        CustomerPipeline.objects.filter(pk=other.pk).update(salesperson=self.other)
        for kind in ('excel', 'pdf'):
            _, content = self.download(kind, {**self.params, 'salesperson': self.person.id})
            if kind == 'excel':
                rows = list(load_workbook(BytesIO(content))['Management Delays'].iter_rows(min_row=7, values_only=True))
                followups = [row for row in rows if row[6] == 'Overdue']
                self.assertEqual(len(followups), 1)
                self.assertEqual(followups[0][1], 'Inside customer')
                self.assertEqual(followups[0][2], self.person.name)
                self.assertEqual(followups[0][3], 'Overdue follow-up: Call customer')
                self.assertEqual(followups[0][5].date(), date(2026, 9, 20))
            else:
                text = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(content)).pages)
                self.assertIn('Overdue follow-up: Call customer', text)
                self.assertIn('Overdue', text)
                self.assertNotIn('Other rep overdue', text)
                self.assertNotIn('Completed followup', text)
                self.assertNotIn('Closed followup', text)

    @patch('marketing.reports.timezone.localdate', return_value=date(2026, 9, 18))
    def test_management_delays_pending_and_overdue_with_contact_dates(self, mocked_today):
        from .reports import build_management_delays
        self.deal('Pending future', date(2026, 9, 16), 10,
                  next_followup_date=date(2026, 10, 5), followup_status='Pending')
        self.deal('Overdue current', date(2026, 9, 16), 10,
                  next_followup_date=date(2026, 9, 17), followup_status='Overdue')
        self.deal('Completed excluded', self.start, 10,
                  next_followup_date=self.start, followup_status='COMPLETED')
        rows = build_management_delays(self.start, self.end, self.person.id)
        pipeline = {row['customer']: row for row in rows if row.get('source') == 'customer_pipeline'}
        self.assertEqual(pipeline['Pending future']['status'], 'Pending')
        self.assertEqual(pipeline['Pending future']['last_contact_date'], date(2026, 9, 16))
        self.assertEqual(pipeline['Pending future']['next_followup_date'], date(2026, 10, 5))
        self.assertEqual(pipeline['Overdue current']['status'], 'Overdue')
        self.assertNotIn('Completed excluded', pipeline)
        _, content = self.download('excel')
        sheet = load_workbook(BytesIO(content))['Management Delays']
        self.assertEqual(sheet['H6'].value, 'Last Contact Date')
        self.assertEqual(sheet['I6'].value, 'Next Follow-up Date')
        row = next(row for row in sheet.iter_rows(min_row=7, values_only=True) if row[1] == 'Pending future')
        self.assertEqual(row[7].date(), date(2026, 9, 16))
        self.assertEqual(row[8].date(), date(2026, 10, 5))
